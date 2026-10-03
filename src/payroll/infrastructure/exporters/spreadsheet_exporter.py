"""Payroll spreadsheet export adapters (CSV/XLSX) -- the inverse of XlsxPayrollImporter.

`XlsxPayrollImporter.read_rows()` pivots a wide-format file into long-format
DTOs using CONCEPT_MAP. These adapters do the inverse pivot: persisted
periods (long format, one PayrollItemDetailDTO per concept) back into the
same wide-format shape, using `inverted_concept_map()` -- the formal inverse
of CONCEPT_MAP built once in xlsx_importer.py, never a second hand-copied
mapping here.
"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from payroll.application.dto import PayrollPeriodDetailDTO
from payroll.application.ports.spreadsheet_exporter import PayrollSpreadsheetExporter
from payroll.infrastructure.importers.xlsx_importer import (
    COMPUTED_ONLY_CONCEPT_COLUMNS,
    inverted_concept_map,
    money_columns,
    wide_columns,
)

# Locale-agnostic money display: OOXML always stores the raw number with '.'
# internally regardless of locale -- Excel renders the decimal/thousands
# separators per the user's own locale at display time. This format string
# just guarantees 2 decimals + thousands grouping everywhere, instead of
# leaving the display up to each user's default cell formatting.
MONEY_NUMBER_FORMAT = "#,##0.00"


def export_columns() -> list[str]:
    """Return the real export's full column order: wide_columns() + computed-only.

    The blank template uses wide_columns() alone (see build_template_dataframe())
    -- computed concepts like INCOME_TAX only ever exist on an already-computed
    persisted period, never on a blank form meant for a brand-new import.
    """
    return [*wide_columns(), *COMPUTED_ONLY_CONCEPT_COLUMNS.values()]


def _period_to_wide_row(
    detail: PayrollPeriodDetailDTO, concept_to_column: dict[str, str]
) -> dict[str, object]:
    """Flatten one persisted period + its items into one wide-format row dict."""
    row: dict[str, object] = {
        "period_month": detail.period_month,
        "period_year": detail.period_year,
        "employer": detail.employer_name,
        "payment_date": detail.payment_date.isoformat(),
        "worked_days": detail.worked_days,
    }
    for item in detail.items:
        column = concept_to_column.get(item.concept_code) or (
            COMPUTED_ONLY_CONCEPT_COLUMNS.get(item.concept_code)
        )
        if column is None:
            # No CONCEPT_MAP column and not a known computed-only concept --
            # not expected today (see design recommendation Item 1), so this
            # is a visible gap rather than a silently dropped column: it
            # simply never gets written, and the next CONCEPT_MAP-drift test
            # (see tests/unit/infrastructure/test_spreadsheet_exporter.py)
            # is what catches it, per "no silent fallbacks."
            continue
        # Kept as a native Decimal, never str(): stringifying here made every
        # XLSX amount a text cell instead of a real number (Excel then shows
        # it exactly as typed, '.' decimal and all, regardless of the user's
        # own locale, and it doesn't sum). CSV output is unaffected -- pandas
        # still renders a Decimal the same way it rendered our old str().
        row[column] = item.amount_clp
    row["net_pay"] = (
        detail.summary.declared_net_pay_clp
        if detail.summary is not None
        and detail.summary.declared_net_pay_clp is not None
        else None
    )
    return row


def build_export_dataframe(periods: list[PayrollPeriodDetailDTO]) -> pd.DataFrame:
    """Build the wide-format DataFrame for a real bulk export."""
    concept_to_column = inverted_concept_map()
    rows = [_period_to_wide_row(period, concept_to_column) for period in periods]
    return pd.DataFrame(rows, columns=export_columns())


def build_template_dataframe() -> pd.DataFrame:
    """Build the blank template DataFrame: header-only, zero data rows."""
    return pd.DataFrame(columns=wide_columns())


def _to_csv_bytes(dataframe: pd.DataFrame) -> bytes:
    buffer = BytesIO()
    dataframe.to_csv(buffer, index=False)
    return buffer.getvalue()


def _to_xlsx_bytes(dataframe: pd.DataFrame) -> bytes:
    """Write the DataFrame as XLSX with real numeric cells for money columns.

    pandas already writes Decimal values as native numbers (openpyxl coerces
    them), which is the important part -- this second pass only adds a
    consistent 2-decimal number_format on top, purely cosmetic, applied only
    to columns that actually exist in this DataFrame (money_columns() lists
    income_tax/unemployment_insurance too, which the blank template doesn't
    have -- iterating dataframe.columns rather than money_columns() means
    that difference needs no special-casing here).
    """
    buffer = BytesIO()
    dataframe.to_excel(buffer, index=False, engine="openpyxl")
    buffer.seek(0)

    workbook = load_workbook(buffer)
    worksheet = workbook.active
    money_column_names = set(money_columns())
    for col_index, column_name in enumerate(dataframe.columns, start=1):
        if column_name not in money_column_names:
            continue
        column_letter = get_column_letter(col_index)
        for cell in worksheet[column_letter][1:]:  # [0] is the header row
            if cell.value is not None:
                cell.number_format = MONEY_NUMBER_FORMAT

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


@dataclass(frozen=True, slots=True)
class CsvPayrollExporter(PayrollSpreadsheetExporter):
    """CSV adapter for PayrollSpreadsheetExporter."""

    def export(self, periods: list[PayrollPeriodDetailDTO]) -> bytes:
        """Serialize the given periods into CSV bytes."""
        return _to_csv_bytes(build_export_dataframe(periods))

    def export_template(self) -> bytes:
        """Serialize a blank, header-only CSV shell."""
        return _to_csv_bytes(build_template_dataframe())


@dataclass(frozen=True, slots=True)
class XlsxPayrollExporter(PayrollSpreadsheetExporter):
    """XLSX adapter for PayrollSpreadsheetExporter."""

    def export(self, periods: list[PayrollPeriodDetailDTO]) -> bytes:
        """Serialize the given periods into XLSX bytes."""
        return _to_xlsx_bytes(build_export_dataframe(periods))

    def export_template(self) -> bytes:
        """Serialize a blank, header-only XLSX shell."""
        return _to_xlsx_bytes(build_template_dataframe())


EXPORTERS_BY_FORMAT: dict[str, PayrollSpreadsheetExporter] = {
    "csv": CsvPayrollExporter(),
    "xlsx": XlsxPayrollExporter(),
}


def get_exporter_for_format(spreadsheet_format: str) -> PayrollSpreadsheetExporter:
    """Resolve the exporter adapter for a `format=csv|xlsx` query parameter.

    Both adapters are stateless (frozen dataclasses with no fields), so this
    intentionally returns the same shared instance rather than constructing a
    new one per request.
    """
    return EXPORTERS_BY_FORMAT[spreadsheet_format]
